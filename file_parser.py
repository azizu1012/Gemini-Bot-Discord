# file_parser.py
import discord
import io
import chardet
import pandas as pd
import openpyxl
import docx
import PyPDF2 # New import for PDF parsing
from typing import Optional, Dict, Any
from config import logger

# --- CÁC THƯ VIỆT CẦN CÀI ĐẶT ---
# pip install python-docx openpyxl pandas chardet PyPDF2 xlrd

async def parse_attachment(attachment: discord.Attachment) -> Optional[Dict[str, str]]:
    """
    Phân tích file đính kèm của Discord và trả về nội dung dưới dạng text.
    Hỗ trợ: txt, md, json, yaml, ini, bat, reg, docx, xlsx, csv, xls, pdf.
    Đối với các file không được hỗ trợ trực tiếp, sẽ cố gắng đọc nội dung thô dưới dạng văn bản.
    """
    filename = attachment.filename
    file_extension = filename.split('.')[-1].lower()
    
    try:
        file_content_bytes = await attachment.read()
        if not file_content_bytes:
            return None

        content = ""
        parsed_successfully = False

        if file_extension in ['txt', 'md', 'json', 'yaml', 'ini', 'bat', 'reg', 'log', 'cfg', 'conf', 'readme']:
            # Tự động nhận diện encoding
            detected_encoding = chardet.detect(file_content_bytes)['encoding'] or 'utf-8'
            try:
                content = file_content_bytes.decode(detected_encoding)
            except UnicodeDecodeError:
                logger.warning(f"Chardet failed for {filename}, falling back to utf-8 with errors.")
                content = file_content_bytes.decode('utf-8', errors='ignore')
            parsed_successfully = True

        elif file_extension == 'docx':
            with io.BytesIO(file_content_bytes) as f:
                doc = docx.Document(f)
                content = "\n".join([para.text for para in doc.paragraphs if para.text])
            parsed_successfully = True

        elif file_extension == 'xlsx':
            with io.BytesIO(file_content_bytes) as f:
                workbook = openpyxl.load_workbook(f, read_only=True)
                for sheet in workbook:
                    content += f"--- Sheet: {sheet.title} ---\n"
                    for row in sheet.iter_rows(values_only=True):
                        # Loại bỏ các cell None và join lại
                        row_data = [str(cell) for cell in row if cell is not None]
                        if row_data:
                            content += ", ".join(row_data) + "\n"
            parsed_successfully = True

        elif file_extension in ['csv', 'xls']:
            # Dùng Pandas để xử lý CSV, XLS (cần 'xlrd' cho .xls)
            with io.BytesIO(file_content_bytes) as f:
                if file_extension == 'csv':
                    # Cố gắng tự nhận diện encoding cho CSV
                    detected_encoding = chardet.detect(file_content_bytes[:5000])['encoding'] or 'utf-8'
                    try:
                        f.seek(0)
                        df = pd.read_csv(f, encoding=detected_encoding)
                    except UnicodeDecodeError:
                        f.seek(0)
                        df = pd.read_csv(f, encoding='utf-8', errors='ignore')
                else: # xls
                    df = pd.read_excel(f, engine='xlrd' if file_extension == 'xls' else 'openpyxl')
                
                content = df.to_string()
            parsed_successfully = True
        
        elif file_extension == 'pdf':
            with io.BytesIO(file_content_bytes) as f:
                reader = PyPDF2.PdfReader(f)
                for page_num in range(len(reader.pages)):
                    page = reader.pages[page_num]
                    content += page.extract_text() + "\n"
            parsed_successfully = True

        if not parsed_successfully:
            # Fallback: attempt to decode as text for any other file type
            detected_encoding = chardet.detect(file_content_bytes)['encoding'] or 'utf-8'
            try:
                content = file_content_bytes.decode(detected_encoding)
                logger.info(f"Đã đọc nội dung thô của file không hỗ trợ '{filename}' dưới dạng văn bản.")
            except UnicodeDecodeError:
                content = file_content_bytes.decode('utf-8', errors='ignore')
                logger.warning(f"Không thể giải mã '{filename}' bằng '{detected_encoding}', đã thử utf-8 với errors='ignore'.")
            
            # If content is still empty or mostly non-text, it might be a binary file
            if not content.strip() or len(content.strip()) < len(file_content_bytes) / 10: # Heuristic for binary
                logger.info(f"Nội dung thô của file '{filename}' có vẻ không phải văn bản hoặc quá ngắn, bỏ qua.")
                return None
            
            # For fallback, we might want to indicate it's a raw text interpretation
            # and explicitly change the extension to .txt as requested by the user.
            base_name = filename.rsplit('.', 1)[0] if '.' in filename else filename
            filename = f"{base_name}.txt"


        if not content.strip():
            logger.info(f"Nội dung file rỗng sau khi parse: {filename}")
            return None

        # Giới hạn kích thước nội dung để tránh quá tải DB
        MAX_CONTENT_LENGTH = 50000 
        if len(content) > MAX_CONTENT_LENGTH:
            content = content[:MAX_CONTENT_LENGTH] + "\n... (Nội dung quá dài, đã cắt bớt)"
            
        logger.info(f"Đã parse thành công file: {filename}")
        return {"filename": filename, "content": content}

    except Exception as e:
        logger.error(f"Lỗi khi parse file {filename}: {e}")
        return None